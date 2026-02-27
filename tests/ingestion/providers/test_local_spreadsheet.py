"""Tests for LocalSpreadsheetProvider — S0-3."""

import csv
import io

import pytest
from openpyxl import Workbook
from uuid_extensions import uuid7

from src.ingestion.providers.local_spreadsheet import LocalSpreadsheetProvider
from src.models.document import DocumentGraph

VALID_CHECKSUM = "sha256:" + "a" * 64

BOQ_CSV_ROWS = [
    ["Description", "Quantity", "Unit", "Unit Price", "Total"],
    ["Structural Steel", "5000", "tonnes", "3500", "17500000"],
    ["Concrete", "20000", "m3", "450", "9000000"],
]


def _make_csv_bytes(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _make_xlsx_bytes(rows: list[list[str | int | float]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _options(checksum: str = VALID_CHECKSUM):
    from src.ingestion.providers.base import ExtractionOptions
    return ExtractionOptions(doc_checksum=checksum)


class TestLocalSpreadsheetProviderCSV:
    """CSV extraction via provider interface."""

    @pytest.mark.anyio
    async def test_extract_csv_returns_document_graph(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_csv_bytes(BOQ_CSV_ROWS)
        doc_id = uuid7()

        graph = await provider.extract(
            content, "text/csv", doc_id, _options(),
        )

        assert isinstance(graph, DocumentGraph)
        assert graph.document_id == doc_id

    @pytest.mark.anyio
    async def test_csv_has_one_table(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_csv_bytes(BOQ_CSV_ROWS)

        graph = await provider.extract(content, "text/csv", uuid7(), _options())

        assert len(graph.pages) == 1
        assert len(graph.pages[0].tables) == 1

    @pytest.mark.anyio
    async def test_csv_table_has_correct_cell_count(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_csv_bytes(BOQ_CSV_ROWS)

        graph = await provider.extract(content, "text/csv", uuid7(), _options())
        table = graph.pages[0].tables[0]

        # 3 rows x 5 cols = 15 cells
        assert len(table.cells) == 15

    @pytest.mark.anyio
    async def test_csv_confidence_is_one(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_csv_bytes(BOQ_CSV_ROWS)

        graph = await provider.extract(content, "text/csv", uuid7(), _options())
        for cell in graph.pages[0].tables[0].cells:
            assert cell.confidence == 1.0


class TestLocalSpreadsheetProviderExcel:
    """Excel extraction via provider interface."""

    @pytest.mark.anyio
    async def test_extract_xlsx_returns_document_graph(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_xlsx_bytes([
            ["Description", "Quantity"],
            ["Steel", 5000],
        ])

        graph = await provider.extract(
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            uuid7(),
            _options(),
        )

        assert isinstance(graph, DocumentGraph)

    @pytest.mark.anyio
    async def test_xlsx_has_one_table(self) -> None:
        provider = LocalSpreadsheetProvider()
        content = _make_xlsx_bytes([
            ["Description", "Quantity"],
            ["Steel", 5000],
        ])

        graph = await provider.extract(
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            uuid7(),
            _options(),
        )

        assert len(graph.pages) == 1
        assert len(graph.pages[0].tables) == 1


class TestLocalSpreadsheetProviderMeta:
    """Provider metadata and interface compliance."""

    def test_name(self) -> None:
        assert LocalSpreadsheetProvider().name == "local-spreadsheet"

    def test_supported_mime_types(self) -> None:
        mimes = LocalSpreadsheetProvider().supported_mime_types()
        assert "text/csv" in mimes
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in mimes
